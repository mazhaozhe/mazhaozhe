"""
tests/test_documents_tools.py —— 文档生成工具测试
==================================================
验证 docx / xlsx / pptx 三个生成器的正常路径与错误路径。
"""

import json

import pytest

from tools.documents import gen_docx, gen_pptx, gen_xlsx, _add_runs
from tools.coding import write_file


# ---------------------------------------------------------------------------
# gen_docx
# ---------------------------------------------------------------------------
def test_gen_docx(tmp_path):
    """Markdown → Word：标题/列表/粗体都能生成。"""
    md = tmp_path / "report.md"
    write_file.invoke({"path": str(md), "content": (
        "# 周报\n\n## 进展\n\n- 完成了**核心模块**\n- 修复了 3 个 Bug\n"
    )})
    out = tmp_path / "report.docx"
    result = gen_docx.invoke({"markdown_path": str(md), "output_path": str(out)})
    assert "已生成" in result and out.exists()
    # 验证内容确实写进去了
    from docx import Document
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert any("周报" in t for t in texts)
    assert any("核心模块" in t for t in texts)


def test_gen_docx_missing_source(tmp_path):
    """源文件不存在 → FileNotFoundError。"""
    with pytest.raises(FileNotFoundError):
        gen_docx.invoke({"markdown_path": str(tmp_path / "nope.md"),
                         "output_path": str(tmp_path / "out.docx")})


# ---------------------------------------------------------------------------
# gen_xlsx
# ---------------------------------------------------------------------------
def test_gen_xlsx(tmp_path):
    """JSON 二维数组 → Excel，含表头。"""
    rows = json.dumps([["姓名", "分数"], ["张三", 90], ["李四", 85]],
                      ensure_ascii=False)
    out = tmp_path / "score.xlsx"
    result = gen_xlsx.invoke({"output_path": str(out), "rows": rows})
    assert "已生成" in result and out.exists()
    from openpyxl import load_workbook
    ws = load_workbook(str(out)).active
    assert ws["A1"].value == "姓名" and ws["B2"].value == 90


def test_gen_xlsx_bad_json(tmp_path):
    """非法 JSON → ValueError 且错误信息里带示例。"""
    with pytest.raises(ValueError, match="示例"):
        gen_xlsx.invoke({"output_path": str(tmp_path / "x.xlsx"),
                         "rows": "{not json"})


# ---------------------------------------------------------------------------
# gen_pptx
# ---------------------------------------------------------------------------
def test_gen_pptx(tmp_path):
    """幻灯片列表 → PPT，封面 + 内容页。"""
    slides = json.dumps([
        {"title": "背景", "content": "项目启动\n需求明确"},
        {"title": "进展", "content": "完成核心模块"},
    ], ensure_ascii=False)
    out = tmp_path / "deck.pptx"
    result = gen_pptx.invoke({"output_path": str(out),
                              "title": "项目汇报", "slides": slides})
    assert "已生成" in result and out.exists()
    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 3  # 封面 + 2 页


# ---------------------------------------------------------------------------
# 粗体解析辅助函数
# ---------------------------------------------------------------------------
def test_add_runs_bold():
    """**粗体** 状态机：偶数段普通、奇数段粗体。"""
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    _add_runs(para, "普通**加粗**普通")
    runs = [(r.text, bool(r.bold)) for r in para.runs]
    assert runs == [("普通", False), ("加粗", True), ("普通", False)]
