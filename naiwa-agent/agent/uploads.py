"""
agent/uploads.py —— 客户端附件上传（文件 / 照片）
=================================================
桌面客户端把用户选择的文件先"暂存"到 data/uploads/staging/，发送消息时
再归档到 data/uploads/<session_id>/，并按类型决定注入方式：

    image     照片/图片：模型支持视觉时以 image_url 多模态块注入；
               不支持时走降级提示（vision_unsupported_message），不调模型
    text      文本/代码/配置：读取前 N 个字符作为摘录拼进消息
    document  docx/xlsx/pptx/pdf：尽量抽取文本摘录（依赖缺失则只记路径）
    other     其他二进制：只记录保存路径，交给 Agent 的文件工具处理

设计要点：
    - 暂存与归档分离：用户加了附件又删掉，不会在会话目录里留垃圾；
    - 所有落盘文件名都加 upload_id 前缀，避免同名覆盖与路径穿越；
    - 扩展名白名单 + 大小上限（agent/config.py：upload_max_mb）。
"""

from __future__ import annotations

import base64
import mimetypes
import re
import shutil
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Optional

from agent.config import settings

# ---------------------------------------------------------------------------
# 分类表
# ---------------------------------------------------------------------------
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".heic"}

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".ndjson",
    ".py", ".pyw", ".js", ".mjs", ".ts", ".java", ".kt", ".c", ".h", ".cpp",
    ".cc", ".hpp", ".go", ".rs", ".rb", ".php", ".swift", ".sql", ".sh",
    ".bat", ".ps1", ".html", ".htm", ".xml", ".svg", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".env", ".log", ".css", ".scss",
}

DOC_EXTS = {".pdf", ".docx", ".xlsx", ".pptx"}

# 允许上传的扩展名总表（白名单之外一律拒绝）
ALLOWED_EXTS = IMAGE_EXTS | TEXT_EXTS | DOC_EXTS

MIME_OVERRIDES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tiff": "image/tiff",
    ".heic": "image/heic",
    ".svg": "image/svg+xml",
    ".md": "text/markdown",
    ".jsonl": "application/jsonl",
    ".ps1": "text/plain",
}

# 视觉能力缺失时返回给用户的降级消息（核心短语：模型无法识别）
VISION_UNSUPPORTED_CORE = "模型无法识别该照片"


class UploadError(Exception):
    """附件相关的用户可读错误（消息直接展示给用户）。"""


# ---------------------------------------------------------------------------
# 目录
# ---------------------------------------------------------------------------
def upload_root() -> Path:
    return settings.abs_path(settings.upload_dir)


def staging_dir() -> Path:
    path = upload_root() / "staging"
    path.mkdir(parents=True, exist_ok=True)
    return path


def session_dir(session_id: str) -> Path:
    safe = re.sub(r"[^0-9A-Za-z_-]", "_", str(session_id or "unknown"))[:80] or "unknown"
    path = upload_root() / safe
    path.mkdir(parents=True, exist_ok=True)
    return path


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------
@dataclass
class Upload:
    """一个附件（暂存态或归档态）的完整信息。"""

    upload_id: str
    name: str                  # 原始文件名（展示用）
    size: int                  # 字节数
    mime: str                  # MIME 类型
    path: str                  # 当前磁盘绝对路径
    kind: str = "other"        # image / text / document / other
    is_image: bool = False
    excerpt: str = ""          # 文本/文档摘录（发送时按需填充）
    meta: dict[str, Any] = field(default_factory=dict)

    def to_public(self) -> dict[str, Any]:
        """给客户端/会话文件的安全视图（路径不含密钥，仅本机路径）。"""
        return {
            "uploadId": self.upload_id,
            "name": self.name,
            "size": self.size,
            "mime": self.mime,
            "kind": self.kind,
            "isImage": self.is_image,
            "path": self.path,
        }


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------
def format_size(size: int) -> str:
    """字节数格式化为人类可读形式（前端 chips 与消息摘要共用）。"""
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / 1024 / 1024:.1f} MB"


def _safe_name(name: str) -> str:
    """去掉路径成分与危险字符，限制长度；保留中文等可见字符。"""
    name = Path(str(name or "file")).name  # 防路径穿越：只取文件名
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", name).strip(" .")
    return (name or "file")[:80]


def _mime_of(name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext in MIME_OVERRIDES:
        return MIME_OVERRIDES[ext]
    guessed, _ = mimetypes.guess_type(name)
    return guessed or "application/octet-stream"


def classify(name: str) -> tuple[str, bool]:
    """按扩展名分类，返回 (kind, is_image)。"""
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXTS:
        return "image", True
    if ext in TEXT_EXTS:
        return "text", False
    if ext in DOC_EXTS:
        return "document", False
    return "other", False


def _validate(filename: str, size: int) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise UploadError(
            f"暂不支持上传 {ext or '无扩展名'} 文件。"
            f"目前支持：图片、文本/代码、docx/xlsx/pptx/pdf。"
        )
    if size <= 0:
        raise UploadError(f"附件 {filename} 是空文件，无法上传。")
    if size > settings.upload_max_bytes:
        raise UploadError(
            f"附件 {filename} 大小 {format_size(size)}，超过上限 "
            f"{settings.upload_max_mb} MB，请压缩或拆分后重试。"
        )


# ---------------------------------------------------------------------------
# 暂存 / 丢弃 / 归档
# ---------------------------------------------------------------------------
def stage_upload(path: Optional[str] = None, data_base64: str = "",
                 name: str = "", mime: str = "") -> Upload:
    """
    把一个文件暂存到 staging 目录。

    两种输入方式（Renderer 优先传 path；拿不到路径时回退 base64）：
        path         本机绝对路径（Electron File.path / webUtils.getPathForFile）
        data_base64  base64 内容（Renderer FileReader 兜底）

    返回 Upload；校验失败抛 UploadError（消息直接给用户看）。
    """
    upload_id = uuid.uuid4().hex[:12]
    target_dir = staging_dir()

    if path:
        src = Path(path)
        if not src.is_file():
            raise UploadError(f"找不到文件：{path}")
        filename = name or src.name
        size = src.stat().st_size
        _validate(filename, size)
        target = target_dir / f"{upload_id}__{_safe_name(filename)}"
        shutil.copy2(src, target)      # 复制而不是移动：用户原文件保持不动
    elif data_base64:
        filename = name or "upload.bin"
        try:
            raw = base64.b64decode(data_base64, validate=False)
        except Exception as exc:
            raise UploadError(f"附件内容解码失败：{exc}") from exc
        size = len(raw)
        _validate(filename, size)
        target = target_dir / f"{upload_id}__{_safe_name(filename)}"
        target.write_bytes(raw)
    else:
        raise UploadError("附件内容为空：既没有文件路径也没有数据。")

    kind, is_image = classify(filename)
    return Upload(
        upload_id=upload_id,
        name=filename,
        size=size,
        mime=mime or _mime_of(filename),
        path=str(target),
        kind=kind,
        is_image=is_image,
    )


def _find_staged(upload_id: str) -> Optional[Path]:
    prefix = f"{re.sub(r'[^0-9A-Za-z_-]', '_', str(upload_id))}__"
    for candidate in staging_dir().glob(f"{prefix}*"):
        if candidate.is_file():
            return candidate
    return None


def discard_upload(upload_id: str) -> bool:
    """删除暂存文件（用户在发送前移除附件时调用）。已归档的不受影响。"""
    staged = _find_staged(upload_id)
    if not staged:
        return False
    try:
        staged.unlink()
        return True
    except OSError:
        return False


def resolve_attachments(items: Iterable[dict[str, Any]],
                        session_id: str) -> list[Upload]:
    """
    把一批暂存附件归档进会话目录，并读好摘录。

    items: [{"uploadId": ..., "name": ..., "mime": ..., "size": ...}, ...]
    暂存文件缺失（过期/被删）时抛 UploadError，由调用方转成用户可见错误。
    """
    resolved: list[Upload] = []
    archive = session_dir(session_id)
    for item in items or []:
        upload_id = str(item.get("uploadId") or "")
        name = item.get("name") or "file"
        staged = _find_staged(upload_id)
        if not staged:
            raise UploadError(
                f"附件 {name} 已过期，请重新选择后再发送。"
            )
        kind, is_image = classify(name)
        target = archive / staged.name          # 保留 upload_id 前缀，防同名覆盖
        if target.exists():
            target.unlink()
        shutil.move(str(staged), str(target))
        upload = Upload(
            upload_id=upload_id,
            name=name,
            size=item.get("size") or target.stat().st_size,
            mime=item.get("mime") or _mime_of(name),
            path=str(target),
            kind=kind,
            is_image=is_image,
        )
        if kind in ("text", "document"):
            upload.excerpt = read_excerpt(upload)
        resolved.append(upload)
    return resolved


def cleanup_staging(max_age_hours: int = 24) -> int:
    """清理超过 max_age_hours 的暂存文件（客户端异常退出时的兜底）。"""
    cutoff = datetime.now() - timedelta(hours=max_age_hours)
    removed = 0
    for candidate in staging_dir().glob("*"):
        try:
            if candidate.is_file() and \
                    datetime.fromtimestamp(candidate.stat().st_mtime) < cutoff:
                candidate.unlink()
                removed += 1
        except OSError:
            continue
    return removed


# ---------------------------------------------------------------------------
# 内容摘录
# ---------------------------------------------------------------------------
def _truncate(text: str, limit: int) -> str:
    text = text.replace("\x00", "")
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n…（仅展示前 {limit} 字符，完整文件见路径）"


def read_excerpt(upload: Upload, limit: Optional[int] = None) -> str:
    """读取文本/文档类附件的摘录；失败返回空串（调用方退化为只记路径）。"""
    limit = limit or settings.upload_text_chars
    path = Path(upload.path)
    try:
        if upload.kind == "text":
            return _truncate(path.read_text(encoding="utf-8", errors="replace"), limit)
        if path.suffix.lower() == ".docx":
            return _truncate(_excerpt_docx(path), limit)
        if path.suffix.lower() == ".xlsx":
            return _truncate(_excerpt_xlsx(path), limit)
        if path.suffix.lower() == ".pptx":
            return _truncate(_excerpt_pptx(path), limit)
        if path.suffix.lower() == ".pdf":
            return _truncate(_excerpt_pdf(path), limit)
    except Exception:
        return ""
    return ""


def _excerpt_docx(path: Path) -> str:
    import docx  # python-docx（项目已依赖）

    lines = [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]
    return "\n".join(lines)


def _excerpt_xlsx(path: Path) -> str:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in book.worksheets[:5]:
        lines.append(f"【工作表】{sheet.title}")
        for row in sheet.iter_rows(max_row=60, values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    book.close()
    return "\n".join(lines)


def _excerpt_pptx(path: Path) -> str:
    from pptx import Presentation

    deck = Presentation(str(path))
    lines: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes
                 if shape.has_text_frame and shape.text_frame.text.strip()]
        if texts:
            lines.append(f"【第 {index} 页】" + " / ".join(texts))
    return "\n".join(lines)


def _excerpt_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""       # 未安装 pypdf：退化为只记路径
    reader = PdfReader(str(path))
    lines = [(page.extract_text() or "") for page in reader.pages[:30]]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 组装给模型 / 给人看的内容
# ---------------------------------------------------------------------------
DEFAULT_IMAGE_PROMPT = "请查看附带的照片并进行分析。"


def build_user_content(text: str, uploads: list[Upload],
                       limit: Optional[int] = None) -> Any:
    """
    组装真正喂给模型的 user 消息内容：

    - 无图片：返回纯文本（附件摘录以【附件：…】块拼在后面）；
    - 有图片：返回 OpenAI 风格多模态块列表（text + image_url data URL）。

    文本为空且有图片时补一句默认提示，避免模型收到"只有图没有说话"。
    """
    limit = limit or settings.upload_text_chars
    blocks: list[dict[str, Any]] = []
    text_blocks: list[str] = []

    for upload in uploads:
        if upload.kind in ("text", "document") and upload.excerpt:
            text_blocks.append(
                f"【附件：{upload.name}（{format_size(upload.size)}，{upload.kind}）】\n"
                f"{upload.excerpt}"
            )
        elif upload.kind == "document":
            text_blocks.append(
                f"【附件：{upload.name}（{format_size(upload.size)}）已保存到 "
                f"{upload.path}，可用文件工具读取】"
            )
        elif upload.kind == "other":
            text_blocks.append(
                f"【附件：{upload.name}（{format_size(upload.size)}）已保存到 "
                f"{upload.path}，可用文件工具读取】"
            )

    full_text = "\n\n".join([t for t in [text.strip(), *text_blocks] if t])
    images = [u for u in uploads if u.is_image]

    if not images:
        return full_text or (text or "").strip()

    prompt = full_text or DEFAULT_IMAGE_PROMPT
    blocks.append({"type": "text", "text": prompt})
    for upload in images:
        mime = upload.mime or _mime_of(upload.name)
        try:
            data = base64.b64encode(Path(upload.path).read_bytes()).decode("ascii")
        except OSError as exc:
            raise UploadError(f"读取图片失败：{upload.name}（{exc}）") from exc
        blocks.append({
            "type": "image_url",
            "image_url": {"url": f"data:{mime};base64,{data}"},
        })
    return blocks


def build_message_display(text: str, uploads: list[Upload]) -> str:
    """会话文件里保存的用户消息文本：原文 + 附件摘要行（不含文件内容）。"""
    lines = []
    for upload in uploads:
        icon = "图片" if upload.is_image else "文件"
        lines.append(f"【{icon}：{upload.name}（{format_size(upload.size)}）】")
    summary = "\n".join(lines)
    body = (text or "").strip()
    if summary and body:
        return f"{body}\n\n{summary}"
    return body or summary


def vision_unsupported_message(model_label: str,
                               images: list[Upload]) -> str:
    """
    当前模型不支持图片输入时的降级回答。

    核心短语固定为"模型无法识别该照片"（产品约定），其余信息帮助用户自救。
    """
    names = "、".join(u.name for u in images[:3])
    if len(images) > 3:
        names += f" 等 {len(images)} 张"
    return (
        f"当前模型（{model_label}）暂不支持图片识别：{VISION_UNSUPPORTED_CORE}"
        f"（{names}）。\n"
        f"照片已保存到本地上传目录。可切换到支持视觉的模型"
        f"（如 qwen-vl-plus、gpt-4o、gemini 等）后重试，"
        f"或把图片里的文字描述给我。"
    )
