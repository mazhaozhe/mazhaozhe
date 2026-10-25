"""
tests/test_uploads.py —— 附件上传模块的单元测试
================================================
覆盖 agent/uploads.py 的核心逻辑：
- classify / format_size / _safe_name：分类与展示
- _validate：扩展名白名单 + 大小上限
- stage_upload / discard_upload / resolve_attachments：暂存 → 归档生命周期
- read_excerpt：文本与 docx/xlsx 摘录
- build_user_content：纯文本 / 多模态块 / 空文本默认提示
- build_message_display：会话里保存的摘要文本
- vision_unsupported_message：降级消息必须包含"模型无法识别"
- _guess_vision：视觉能力猜测
"""

from pathlib import Path

import pytest

from agent import uploads
from agent.config import settings
from agent.models import ModelSpec, _guess_vision

PNG_BYTES = bytes([137, 80, 78, 71, 13, 10, 26, 10]) + b"fake-png-body"


@pytest.fixture(autouse=True)
def isolated_upload_dir(tmp_path, monkeypatch):
    """把上传目录指到临时目录，避免污染项目 data/。"""
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    yield tmp_path / "uploads"


def make_file(directory: Path, name: str, content: bytes = b"hello") -> Path:
    path = directory / name
    path.write_bytes(content.encode() if isinstance(content, str) else content)
    return path


# ---------------------------------------------------------------------------
# 分类 / 展示 / 安全文件名
# ---------------------------------------------------------------------------
class TestClassify:
    def test_image_extensions(self):
        for name in ("a.png", "a.jpg", "a.webp", "a.HEIC"):
            assert uploads.classify(name) == ("image", True)

    def test_svg_is_text_not_image(self):
        # SVG 是 XML 文本，视觉 API 不认，按文本处理
        assert uploads.classify("a.svg") == ("text", False)

    def test_text_extensions(self):
        assert uploads.classify("main.py") == ("text", False)
        assert uploads.classify("notes.md") == ("text", False)
        assert uploads.classify("data.csv") == ("text", False)

    def test_document_extensions(self):
        assert uploads.classify("report.pdf") == ("document", False)
        assert uploads.classify("table.xlsx") == ("document", False)

    def test_unknown_extension(self):
        assert uploads.classify("archive.rar") == ("other", False)


class TestFormatSize:
    def test_bytes(self):
        assert uploads.format_size(512) == "512 B"

    def test_kilobytes(self):
        assert uploads.format_size(2048) == "2.0 KB"

    def test_megabytes(self):
        assert uploads.format_size(3 * 1024 * 1024) == "3.0 MB"


class TestSafeName:
    def test_strips_path_components(self):
        assert uploads._safe_name("../../etc/passwd") == "passwd"

    def test_replaces_dangerous_chars(self):
        assert ":" not in uploads._safe_name('a:b*c?"d.txt')

    def test_keeps_chinese_name(self):
        assert uploads._safe_name("会议纪要.txt") == "会议纪要.txt"

    def test_empty_fallback(self):
        assert uploads._safe_name("") == "file"


# ---------------------------------------------------------------------------
# 校验
# ---------------------------------------------------------------------------
class TestValidate:
    def test_rejects_unknown_extension(self):
        with pytest.raises(uploads.UploadError) as exc:
            uploads._validate("virus.exe", 1024)
        assert "暂不支持上传" in str(exc.value)

    def test_rejects_empty_file(self):
        with pytest.raises(uploads.UploadError) as exc:
            uploads._validate("a.png", 0)
        assert "空文件" in str(exc.value)

    def test_rejects_oversize(self, monkeypatch):
        monkeypatch.setattr(settings, "upload_max_mb", 1)
        with pytest.raises(uploads.UploadError) as exc:
            uploads._validate("big.png", 2 * 1024 * 1024)
        assert "超过上限" in str(exc.value)

    def test_accepts_allowed(self):
        uploads._validate("photo.png", 1024)   # 不抛异常即通过


# ---------------------------------------------------------------------------
# 暂存 / 丢弃 / 归档
# ---------------------------------------------------------------------------
class TestStageUpload:
    def test_stage_by_path(self, tmp_path):
        src = make_file(tmp_path, "photo.png", PNG_BYTES)
        upload = uploads.stage_upload(path=str(src))
        assert upload.is_image and upload.kind == "image"
        assert upload.name == "photo.png"
        assert upload.size == len(PNG_BYTES)
        assert Path(upload.path).is_file()
        assert src.is_file()                      # 原文件不动（复制而非移动）
        assert upload.path.startswith(str(uploads.staging_dir()))

    def test_stage_by_base64(self):
        import base64
        upload = uploads.stage_upload(
            data_base64=base64.b64encode(PNG_BYTES).decode(), name="shot.png")
        assert upload.is_image
        assert Path(upload.path).read_bytes() == PNG_BYTES

    def test_stage_missing_path(self, tmp_path):
        with pytest.raises(uploads.UploadError):
            uploads.stage_upload(path=str(tmp_path / "nope.png"))

    def test_stage_without_any_input(self):
        with pytest.raises(uploads.UploadError):
            uploads.stage_upload()

    def test_discard_staged(self, tmp_path):
        src = make_file(tmp_path, "a.txt")
        upload = uploads.stage_upload(path=str(src))
        assert uploads.discard_upload(upload.upload_id) is True
        assert not Path(upload.path).exists()
        assert uploads.discard_upload(upload.upload_id) is False

    def test_discard_unknown_id(self):
        assert uploads.discard_upload("not-exist") is False


class TestResolveAttachments:
    def _stage(self, tmp_path, name="photo.png"):
        src = make_file(tmp_path, name, PNG_BYTES if name.endswith(".png") else b"text")
        return uploads.stage_upload(path=str(src))

    def test_resolve_moves_into_session_dir(self, tmp_path):
        upload = self._stage(tmp_path)
        resolved = uploads.resolve_attachments(
            [{"uploadId": upload.upload_id, "name": upload.name}], "client-1")
        assert len(resolved) == 1
        target = Path(resolved[0].path)
        assert target.is_file()
        assert target.parent == uploads.session_dir("client-1")
        assert not list(uploads.staging_dir().glob(f"{upload.upload_id}__*"))  # 暂存已清空

    def test_resolve_missing_staged_raises(self):
        with pytest.raises(uploads.UploadError) as exc:
            uploads.resolve_attachments(
                [{"uploadId": "gone", "name": "a.png"}], "client-1")
        assert "已过期" in str(exc.value)

    def test_resolve_reads_text_excerpt(self, tmp_path):
        src = make_file(tmp_path, "notes.txt", "第一行\n第二行")
        upload = uploads.stage_upload(path=str(src))
        resolved = uploads.resolve_attachments(
            [{"uploadId": upload.upload_id, "name": "notes.txt"}], "client-1")
        assert resolved[0].excerpt == "第一行\n第二行"


# ---------------------------------------------------------------------------
# 摘录
# ---------------------------------------------------------------------------
class TestExcerpt:
    def test_text_excerpt(self, tmp_path):
        src = make_file(tmp_path, "a.txt", "x" * 100)
        upload = uploads.stage_upload(path=str(src))
        excerpt = uploads.read_excerpt(upload, limit=10)
        assert excerpt.startswith("x" * 10)
        assert "仅展示前 10 字符" in excerpt

    def test_docx_excerpt(self, tmp_path):
        import docx
        path = tmp_path / "doc.docx"
        doc = docx.Document()
        doc.add_paragraph("标题一")
        doc.add_paragraph("正文内容")
        doc.save(str(path))
        upload = uploads.stage_upload(path=str(path))
        excerpt = uploads.read_excerpt(upload)
        assert "标题一" in excerpt and "正文内容" in excerpt

    def test_xlsx_excerpt(self, tmp_path):
        import openpyxl
        path = tmp_path / "sheet.xlsx"
        book = openpyxl.Workbook()
        book.active["A1"] = "姓名"
        book.active["B1"] = "分数"
        book.save(str(path))
        upload = uploads.stage_upload(path=str(path))
        assert "姓名" in uploads.read_excerpt(upload)

    def test_unknown_kind_returns_empty(self, tmp_path):
        src = make_file(tmp_path, "a.txt")
        upload = uploads.stage_upload(path=str(src))
        upload.kind = "other"
        assert uploads.read_excerpt(upload) == ""


# ---------------------------------------------------------------------------
# 组装模型输入 / 展示文本
# ---------------------------------------------------------------------------
def fake_upload(upload_id="u1", name="photo.png", is_image=True, path="/tmp/x.png",
                excerpt="", kind="image"):
    return uploads.Upload(upload_id=upload_id, name=name, size=2048,
                          mime="image/png", path=path, kind=kind, is_image=is_image,
                          excerpt=excerpt)


class TestBuildUserContent:
    def test_text_only_returns_string(self):
        content = uploads.build_user_content("你好", [])
        assert content == "你好"

    def test_text_with_text_attachment_appends_excerpt(self):
        upload = fake_upload(name="notes.txt", is_image=False,
                             kind="text", excerpt="摘录内容")
        content = uploads.build_user_content("总结一下", [upload])
        assert content.startswith("总结一下")
        assert "【附件：notes.txt" in content
        assert "摘录内容" in content

    def test_other_file_note(self):
        upload = fake_upload(name="data.bin", is_image=False, kind="other",
                             path="C:/x/data.bin")
        content = uploads.build_user_content("", [upload])
        assert "已保存到 C:/x/data.bin" in content

    def test_image_returns_multimodal_blocks(self, tmp_path):
        import base64
        path = tmp_path / "photo.png"
        path.write_bytes(PNG_BYTES)
        upload = fake_upload(path=str(path))
        content = uploads.build_user_content("这是什么", [upload])
        assert isinstance(content, list)
        assert content[0] == {"type": "text", "text": "这是什么"}
        assert content[1]["type"] == "image_url"
        assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")
        decoded = base64.b64decode(content[1]["image_url"]["url"].split(",", 1)[1])
        assert decoded == PNG_BYTES

    def test_image_without_text_uses_default_prompt(self, tmp_path):
        path = tmp_path / "photo.png"
        path.write_bytes(PNG_BYTES)
        upload = fake_upload(path=str(path))
        content = uploads.build_user_content("", [upload])
        assert content[0]["text"] == uploads.DEFAULT_IMAGE_PROMPT


class TestBuildMessageDisplay:
    def test_display_with_attachments(self):
        uploads_list = [fake_upload(name="a.png"), fake_upload(upload_id="u2", name="b.txt", is_image=False)]
        display = uploads.build_message_display("看这个", uploads_list)
        assert display.startswith("看这个")
        assert "【图片：a.png（2.0 KB）】" in display
        assert "【文件：b.txt（2.0 KB）】" in display

    def test_display_attachment_only(self):
        display = uploads.build_message_display("", [fake_upload()])
        assert display.startswith("【图片：")


# ---------------------------------------------------------------------------
# 视觉降级消息
# ---------------------------------------------------------------------------
class TestVisionUnsupportedMessage:
    def test_core_phrase_present(self):
        message = uploads.vision_unsupported_message("deepseek-chat", [fake_upload()])
        assert "模型无法识别该照片" in message
        assert "deepseek-chat" in message

    def test_multiple_images_truncated_names(self):
        images = [fake_upload(upload_id=f"u{i}", name=f"p{i}.png") for i in range(5)]
        message = uploads.vision_unsupported_message("kimi-k2", images)
        assert "等 5 张" in message


# ---------------------------------------------------------------------------
# 视觉能力猜测
# ---------------------------------------------------------------------------
class TestGuessVision:
    def test_vision_models(self):
        assert _guess_vision("qwen-vl-plus")
        assert _guess_vision("qwen2.5-vl-72b-instruct")
        assert _guess_vision("moonshot-v1-128k-vision-preview")
        assert _guess_vision("gpt-4o-mini")
        assert _guess_vision("gemini-1.5-pro")

    def test_text_models(self):
        assert not _guess_vision("deepseek-chat")
        assert not _guess_vision("deepseek-reasoner")
        assert not _guess_vision("qwen-plus")
        assert not _guess_vision("kimi-k2-0905-preview")

    def test_explicit_patterns_override(self):
        assert _guess_vision("my-custom-eye", patterns=["eye"])
        assert not _guess_vision("qwen-vl-plus", patterns=["nomatch"])


class TestModelSpecVision:
    def test_from_dict_reads_supports_vision(self):
        spec = ModelSpec.from_dict({
            "id": "qwen-vl-plus", "label": "x", "provider": "openai-compatible",
            "model": "qwen-vl-plus", "supports_vision": True,
        })
        assert spec.supports_vision is True

    def test_default_is_false(self):
        spec = ModelSpec.from_dict({"id": "a", "label": "x", "provider": "p", "model": "m"})
        assert spec.supports_vision is False

    def test_public_dict_exposes_vision(self):
        spec = ModelSpec.from_dict({
            "id": "a", "label": "x", "provider": "p", "model": "m",
            "supports_vision": True,
        })
        assert spec.to_public_dict(available=True, current=True)["supports_vision"] is True
